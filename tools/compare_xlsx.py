# Compares two exports as Excel shows them: sheet names, every cell's value and number format,
# bold headers, column widths, frozen panes and tables. Prints only where they differ, never
# values (the exports hold real feed data).
#
#   uv run python -m tools.compare_xlsx <a.xlsx> <b.xlsx> [skip-cell ...]   e.g. About!B1
import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def describe(path: str, skip: set[str]) -> dict:
    wb = load_workbook(path)
    found = {"sheets": wb.sheetnames}
    for ws in wb.worksheets:
        cells = {}
        for row in ws.iter_rows():
            for cell in row:
                if f"{ws.title}!{cell.coordinate}" in skip or cell.value is None:
                    continue
                # ClosedXML's default format reads back as "", openpyxl's as "General": the same.
                fmt = cell.number_format or "General"
                cells[cell.coordinate] = (cell.value, fmt, bool(cell.font and cell.font.b))
        found[ws.title] = {
            "cells": cells,
            # One <col> element can cover a run of columns (ClosedXML groups equal widths).
            "widths": {
                get_column_letter(i): round(v.width, 2)
                for v in ws.column_dimensions.values()
                if v.width
                for i in range(v.min, v.max + 1)
            },
            "freeze": ws.freeze_panes,
            "tables": sorted(
                (t.displayName, t.ref, t.tableStyleInfo.name if t.tableStyleInfo else None)
                for t in ws.tables.values()
            ),
        }
    return found


def main(a: str, b: str, *skip: str) -> int:
    left, right = describe(a, set(skip)), describe(b, set(skip))
    problems = []
    if left["sheets"] != right["sheets"]:
        problems.append("sheet names")
    empty = {"cells": {}, "widths": {}, "freeze": None, "tables": []}
    for sheet in left["sheets"]:
        x, y = left[sheet], right.get(sheet, empty)
        for key in ("widths", "freeze", "tables"):
            if x[key] != y[key]:
                problems.append(f"{sheet}: {key}")
        both = set(x["cells"]) & set(y["cells"])
        bad = sorted(set(x["cells"]) ^ set(y["cells"]))
        bad += sorted(c for c in both if x["cells"][c] != y["cells"][c])
        if bad:
            problems.append(f"{sheet}: {len(bad)} cells differ, first {', '.join(bad[:5])}")
    print("identical" if not problems else "\n".join(problems))
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main(*sys.argv[1:]))
