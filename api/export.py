import io
from dataclasses import dataclass
from datetime import date, datetime, tzinfo

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from feed.availability import TechDay
from feed.diagnostics import Check
from feed.outcomes import CHECKPOINTS, OUTCOMES, REACHED, DayOutcome
from feed.parse import Block

DATE = "yyyy-mm-dd"
TIME = "yyyy-mm-dd hh:mm"
HOURS = "0.00"
PERCENT = "0.0%"

DEFINITIONS = (
    ("Shift h", "Scheduled shift hours."),
    ("Available h", "Shift minus lunch (12-1, Sat 1-2) minus time off."),
    (
        "Booked h",
        "Jobs and tickets inside available time. Canceled/unnecessary tasks and "
        "closed/deleted/held/cleared tickets don't count.",
    ),
    (
        "Free h",
        "Open slots of at least 60 minutes; today's start no sooner than 30 minutes "
        "after the snapshot was read.",
    ),
    ("Utilization", "Booked h / available h."),
    (
        "Unassigned h",
        "Scheduled jobs and tickets with no tech yet, in the region their address maps to.",
    ),
    ("Net h", "Free h minus unassigned h: capacity left once unassigned work is placed."),
    ("Techs off", "Techs with time off overlapping their shift, or a day off entirely."),
    (
        "Planned",
        "Install jobs and FIELD/TC tickets assigned for the day in the first snapshot "
        "between 06:00 and 07:00.",
    ),
    (
        "d0 / d1 / d2",
        "The last snapshot before midnight ending the plan day, the next day, "
        "and the day after.",
    ),
    ("Outcome", "The latest checkpoint's result: d2 for final days, so far for provisional."),
    (
        "Pulled d0",
        "Planned jobs canceled, unscheduled or rescheduled by the end of the plan day, "
        "split by how far the tech got (in progress, en route).",
    ),
    ("Provisional", "D+2 hasn't ended yet, so the day's outcomes can still change."),
)


@dataclass
class Column:
    header: str
    fmt: str | None = None
    width: int | None = None


def write_sheet(wb: Workbook, title: str, columns: list[Column], rows: list[list]) -> None:
    ws = wb.create_sheet(title)
    ws.append([c.header for c in columns])
    for row in rows:
        ws.append(row)
    for index, column in enumerate(columns, start=1):
        letter = get_column_letter(index)
        if column.fmt:
            for (cell,) in ws.iter_rows(min_row=2, min_col=index, max_col=index):
                cell.number_format = column.fmt
        longest = max([len(column.header)] + [len(str(row[index - 1] or "")) for row in rows[:500]])
        ws.column_dimensions[letter].width = column.width or min(max(longest + 2, 8), 50)
    ws.freeze_panes = "A2"
    end = f"{get_column_letter(len(columns))}{len(rows) + 1}"
    if rows:
        # Excel rejects a table with no data rows; an empty sheet keeps just its header.
        table = Table(displayName=title.replace(" ", ""), ref=f"A1:{end}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
    else:
        for cell in ws[1]:
            cell.font = Font(bold=True)


CAPACITY_COLUMNS = [
    Column("Techs on"),
    Column("Techs off"),
    Column("Shift h", HOURS),
    Column("Available h", HOURS),
    Column("Booked h", HOURS),
    Column("Free h", HOURS),
    Column("Utilization", PERCENT),
    Column("Jobs"),
    Column("Tickets"),
    Column("Unassigned jobs"),
    Column("Unassigned tickets"),
    Column("Unassigned h", HOURS),
    Column("Net h", HOURS),
]
CAPACITY_FIELDS = [
    "techs_on",
    "techs_off",
    "shift_h",
    "available_h",
    "booked_h",
    "free_h",
    "utilization",
    "jobs",
    "tickets",
    "unassigned_jobs",
    "unassigned_tickets",
    "unassigned_h",
    "net_h",
]


def summary_rows(entries: list[dict]) -> list[list]:
    rows = []
    for entry in entries:
        rows.append([entry["date"], "(all)"] + [entry["totals"][f] for f in CAPACITY_FIELDS])
        for region in entry["by_region"]:
            rows.append(
                [entry["date"], region["region"] or "(none)"] + [region[f] for f in CAPACITY_FIELDS]
            )
    return rows


OUTCOME_DAY_COLUMNS = (
    [
        Column("Date", DATE, 12),
        Column("Status"),
        Column("Provisional"),
        Column("Kind"),
        Column("Planned"),
    ]
    + [Column(f"Completed {c}") for c in CHECKPOINTS]
    + [Column(f"Completion {c}", PERCENT) for c in CHECKPOINTS]
    + [Column(o.replace("_", " ").capitalize()) for o in OUTCOMES]
    + [Column("Pulled d0")]
    + [Column(f"Pulled {r.replace('_', ' ')}") for r in REACHED]
    + [Column("Pulled, prereqs open")]
)


def outcome_day_rows(kpis: dict) -> list[list]:
    rows = []
    for entry in kpis["days"]:
        if entry["status"] != "ok":
            row = [entry["date"], entry["status"], entry["provisional"]]
            rows.append(row + [None] * (len(OUTCOME_DAY_COLUMNS) - len(row)))
            continue
        for kind in ("job", "ticket"):
            stats = entry[kind]
            pulled = stats.get("pulled_d0", {})
            rows.append(
                [entry["date"], entry["status"], entry["provisional"], kind, stats["planned"]]
                + [stats["completed"][c] for c in CHECKPOINTS]
                + [stats["completion_rate"][c] for c in CHECKPOINTS]
                + [stats["outcome"][o] for o in OUTCOMES]
                + [pulled.get("total")]
                + [pulled.get(r) for r in REACHED]
                + [pulled.get("prereqs_open")]
            )
    return rows


DIAGNOSTIC_COLUMNS = [
    Column("Check"),
    Column("Severity"),
    Column("Ref"),
    Column("Kind"),
    Column("Status"),
    Column("Date", DATE, 12),
    Column("Starts", TIME, 17),
    Column("Ends", TIME, 17),
    Column("Tech id"),
    Column("Tech"),
    Column("Region"),
    Column("Detail", width=60),
]
DIAGNOSTIC_FIELDS = (
    "ref_id",
    "kind",
    "status",
    "work_date",
    "starts_at",
    "ends_at",
    "tech_id",
    "tech_name",
    "region",
)


def diagnostic_rows(checks: list[Check]) -> list[list]:
    # Fixed columns for the job or tech, and whatever else a check reports as "name: value".
    rows = []
    for check in checks:
        for row in check.rows:
            detail = "; ".join(
                f"{name.replace('_', ' ')}: {value}"
                for name, value in row.items()
                if name not in DIAGNOSTIC_FIELDS
            )
            rows.append(
                [check.title, check.severity] + [row.get(f) for f in DIAGNOSTIC_FIELDS] + [detail]
            )
    return rows


def workbook(
    *,
    tz: tzinfo,
    meta: dict,
    filters: dict,
    entries: list[dict],
    days: list[TechDay],
    demand: list[Block],
    schedule: list[Block],
    outcomes: list[tuple[DayOutcome, bool]],
    kpis: dict,
    checks: list[Check],
) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    write_sheet(
        wb,
        "Summary",
        [Column("Date", DATE, 12), Column("Region")] + CAPACITY_COLUMNS,
        summary_rows(entries),
    )
    write_sheet(
        wb,
        "Tech days",
        [
            Column("Date", DATE, 12),
            Column("Tech id"),
            Column("Tech"),
            Column("Region"),
            Column("Skills"),
            Column("Shift h", HOURS),
            Column("Lunch h", HOURS),
            Column("Time off h", HOURS),
            Column("Available h", HOURS),
            Column("Booked h", HOURS),
            Column("Free h", HOURS),
            Column("Jobs"),
            Column("Tickets"),
            Column("On time off"),
        ],
        [
            [
                d.work_date,
                d.tech_id,
                d.tech_name,
                d.region,
                d.skills,
                d.shift_hours,
                d.lunch_hours,
                d.time_off_hours,
                d.available_hours,
                d.booked_hours,
                d.free_hours,
                d.jobs,
                d.tickets,
                d.on_time_off,
            ]
            for d in days
        ],
    )
    write_sheet(
        wb,
        "Free slots",
        [
            Column("Date", DATE, 12),
            Column("Tech id"),
            Column("Tech"),
            Column("Region"),
            Column("Skills"),
            Column("Open from", TIME, 17),
            Column("Open until", TIME, 17),
            Column("Minutes"),
        ],
        [
            [
                s.work_date,
                s.tech_id,
                s.tech_name,
                s.region,
                s.skills,
                s.open_from,
                s.open_until,
                s.open_minutes,
            ]
            for d in days
            for s in d.free
        ],
    )
    block_columns = [
        Column("Date", DATE, 12),
        Column("Kind"),
        Column("Ref"),
        Column("Status"),
        Column("Task type"),
        Column("Tech id"),
        Column("Tech"),
        Column("Region"),
        Column("Starts", TIME, 17),
        Column("Ends", TIME, 17),
        Column("Address issue"),
    ]

    def block_row(b: Block) -> list:
        return [
            b.work_date,
            b.kind,
            b.ref_id,
            b.status,
            b.task_type,
            b.tech_id,
            b.tech_name,
            b.region,
            b.starts_at,
            b.ends_at,
            b.address_issue,
        ]

    write_sheet(wb, "Schedule", block_columns, [block_row(b) for b in schedule])
    write_sheet(wb, "Unassigned work", block_columns, [block_row(b) for b in demand])
    write_sheet(
        wb,
        "Outcomes by day",
        OUTCOME_DAY_COLUMNS,
        outcome_day_rows(kpis),
    )
    region, tech = filters.get("region"), filters.get("tech")
    write_sheet(
        wb,
        "Outcome items",
        [
            Column("Plan date", DATE, 12),
            Column("Provisional"),
            Column("Kind"),
            Column("Ref"),
            Column("Tech id"),
            Column("Tech"),
            Column("Region"),
            Column("Task type"),
            Column("Planned start", TIME, 17),
            Column("d0"),
            Column("d1"),
            Column("d2"),
            Column("Reached"),
            Column("Prereqs open"),
        ],
        [
            [
                outcome.day,
                provisional,
                i.kind,
                i.ref_id,
                i.tech_id,
                i.tech_name,
                i.region,
                i.task_type,
                i.planned_start,
                i.outcomes.get("d0"),
                i.outcomes.get("d1"),
                i.outcomes.get("d2"),
                i.reached,
                i.prereqs_open,
            ]
            for outcome, provisional in sorted(outcomes, key=lambda o: o[0].day)
            for i in outcome.items
            if (region is None or i.region == region) and (tech is None or i.tech_id == tech)
        ],
    )

    write_sheet(wb, "Diagnostics", DIAGNOSTIC_COLUMNS, diagnostic_rows(checks))

    about = wb.create_sheet("About")
    # Excel has no time zones: both are written in MBS local time, like the block timestamps.
    generated = meta.get("generated_at")
    about.append(["Exported at", datetime.now(tz).replace(microsecond=0, tzinfo=None)])
    about.append(["Snapshot id", meta.get("id")])
    about.append(
        ["Snapshot generated at", generated and generated.astimezone(tz).replace(tzinfo=None)]
    )
    for name, value in filters.items():
        about.append([f"Filter: {name}", value if value is not None else "(all)"])
    about.append([])
    about.append(["Metric", "Definition"])
    about.cell(about.max_row, 1).font = Font(bold=True)
    about.cell(about.max_row, 2).font = Font(bold=True)
    for name, text in DEFINITIONS:
        about.append([name, text])
    about.append([])
    about.append(["Diagnostic", "Rows"])
    about.cell(about.max_row, 1).font = Font(bold=True)
    about.cell(about.max_row, 2).font = Font(bold=True)
    for check in checks:
        about.append([check.title, len(check.rows) if check.available else "not in feed yet"])
    for row in about.iter_rows():
        for cell in row:
            if isinstance(cell.value, datetime):
                cell.number_format = TIME
            elif isinstance(cell.value, date):
                cell.number_format = DATE
    about.column_dimensions["A"].width = 24
    about.column_dimensions["B"].width = 100

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
