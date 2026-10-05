import io
from datetime import UTC, date, datetime
from zoneinfo import ZoneInfo

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from api.main import create_app
from feed.parse import Block

TZ = ZoneInfo("America/Chicago")
DAY = date(2026, 10, 6)
NOW = datetime(2026, 10, 6, 12, 0, tzinfo=UTC)  # 07:00 in Chicago
GENERATED = datetime(2026, 10, 6, 11, 5, tzinfo=UTC)  # 06:05: the morning plan


def block(kind, start, end, tech="a", region="", skills="", ref="", status="A") -> Block:
    starts_at = datetime.fromisoformat(f"2026-10-06 {start}")
    return Block(
        kind=kind,
        work_date=DAY,
        tech_id=tech,
        tech_name=tech.upper() if tech else "",
        starts_at=starts_at,
        ends_at=datetime.fromisoformat(f"2026-10-06 {end}"),
        ref_id=ref,
        status=status,
        department="FIELD",
        region=region,
        skills=skills,
        task_type="3" if kind == "job" else "",
    )


BLOCKS = [
    block("shift", "08:00", "17:00", "a", "North", "INS, RECO"),
    block("shift", "08:00", "17:00", "b", "South", "INS"),
    block("job", "08:00", "10:00", "a", ref="j1"),
    block("job_unassigned", "09:00", "11:00", "", "North", ref="u1"),
]
META = {"id": 7, "format": "blocks", "generated_at": GENERATED, "ingested_at": GENERATED}


class FakeStore:
    def __init__(self, blocks=BLOCKS):
        self.blocks = blocks

    def latest(self):
        return {"snapshot": META, "blocks": self.blocks} if self.blocks else None

    def snapshot_meta(self):
        return META if self.blocks else None

    def blocks_snapshots(self):
        return [{"id": 7, "generated_at": GENERATED}] if self.blocks else []

    def work_blocks(self, snapshot_id):
        return [b for b in self.blocks if b.kind not in ("shift", "time_off")]

    def shift_regions(self, snapshot_id, day):
        return {b.tech_id: b.region for b in self.blocks if b.kind == "shift"}

    def load_day(self, day):
        return None


def client(store=None, **kwargs) -> TestClient:
    app = create_app(
        store=store or FakeStore(), tz=TZ, cors_origins=[], clock=lambda: NOW, **kwargs
    )
    return TestClient(app)


@pytest.fixture
def api():
    return client(api_key="")


def test_health_reports_snapshot_age_without_auth():
    body = client(api_key="secret").get("/health").json()
    assert body["ok"] is True
    assert body["snapshot"]["age_min"] == 55
    assert body["snapshot"]["stale"] is True


def test_api_key_is_required_when_set():
    secured = client(api_key="secret")
    assert secured.get("/api/v1/filters").status_code == 401
    assert secured.get("/api/v1/filters", headers={"X-API-Key": "wrong"}).status_code == 401
    assert secured.get("/api/v1/filters", headers={"X-API-Key": "secret"}).status_code == 200


def test_404_before_the_first_snapshot():
    empty = client(FakeStore(blocks=[]), api_key="")
    assert empty.get("/api/v1/calendar").status_code == 404
    assert empty.get("/health").json()["snapshot"] is None


def test_filters(api):
    body = api.get("/api/v1/filters").json()
    assert body["regions"] == ["North", "South"]
    assert body["skills"] == ["INS", "RECO"]
    assert [t["tech_id"] for t in body["techs"]] == ["a", "b"]


def test_calendar_starts_today_in_mbs_time(api):
    body = api.get("/api/v1/calendar", params={"days": 3}).json()
    assert [d["date"] for d in body["days"]] == ["2026-10-06", "2026-10-07", "2026-10-08"]
    today = body["days"][0]
    assert today["totals"]["techs_on"] == 2
    assert today["totals"]["unassigned_jobs"] == 1
    assert [r["region"] for r in today["by_region"]] == ["North", "South"]
    assert body["snapshot"]["id"] == 7


def test_calendar_rejects_too_long_a_range(api):
    assert api.get("/api/v1/calendar", params={"days": 400}).status_code == 422


def test_calendar_day_drill_down(api):
    body = api.get("/api/v1/calendar/2026-10-06", params={"region": "North"}).json()
    [tech] = body["techs"]
    assert tech["tech_id"] == "a"
    assert [w["ref_id"] for w in tech["work"]] == ["j1"]
    # Now is 07:00, so the lead time doesn't clip the 10:00 start.
    assert tech["free"][0]["open_from"] == "2026-10-06T10:00:00"
    assert [u["ref_id"] for u in body["unassigned"]] == ["u1"]
    assert body["totals"]["techs_on"] == 1


def test_capacity_kpis(api):
    body = api.get("/api/v1/kpis/capacity", params={"days": 2}).json()
    assert len(body["series"]) == 2
    assert body["series"][0]["utilization"] == round(2 / 16, 3)
    assert [r["region"] for r in body["by_region"]] == ["North", "South"]


def test_outcome_kpis(api):
    body = api.get("/api/v1/kpis/outcomes", params={"start": "2026-10-06", "end": "2026-10-06"})
    body = body.json()
    [day] = body["days"]
    assert day["provisional"] is True
    assert day["job"]["planned"] == 1
    assert body["by_region"][0]["region"] == "North"


def test_outcome_range_is_validated(api):
    params = {"start": "2026-10-06", "end": "2026-10-01"}
    assert api.get("/api/v1/kpis/outcomes", params=params).status_code == 422


def test_export_is_a_workbook(api):
    response = api.get("/api/v1/export.xlsx", params={"days": 2, "history_days": 1})
    assert response.status_code == 200
    assert "techavail_2026-10-06.xlsx" in response.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(response.content))
    assert wb.sheetnames == [
        "Summary",
        "Tech days",
        "Free slots",
        "Schedule",
        "Unassigned work",
        "Outcomes by day",
        "Outcome items",
        "Diagnostics",
        "About",
    ]
    summary = list(wb["Summary"].iter_rows(values_only=True))
    assert summary[0][:3] == ("Date", "Region", "Techs on")
    assert summary[1][:3] == (datetime(2026, 10, 6), "(all)", 2)
    assert len(list(wb["Schedule"].iter_rows())) == 1 + 3
    items = list(wb["Outcome items"].iter_rows(values_only=True))
    assert items[1][3] == "j1"
    headers = {c.value for ws in wb for c in ws[1] if isinstance(c.value, str)}
    assert not {h for h in headers if "lat" in h.lower() or "lon" in h.lower()}
    assert "Address issue" in headers


def located(b: Block, lat: float, lon: float, issue: str = "") -> Block:
    b.latitude, b.longitude, b.address_issue = lat, lon, issue
    return b


MAPPED = [
    block("shift", "08:00", "17:00", "a", "North", "INS"),
    located(block("job", "08:00", "10:00", "a", ref="j1"), 40.123456, -100.654321),
    located(block("job", "08:00", "10:00", "b", ref="j1"), 40.123456, -100.654321),
    located(block("job_unassigned", "09:00", "11:00", "", "South", ref="u1"), 0, 0, "no_gps"),
    located(block("job", "12:00", "13:00", "a", ref="j2", status="X"), 38.9, -92.3),
]
for b in MAPPED:
    if b.latitude == 0:
        b.latitude = b.longitude = None


def test_map_rounds_coordinates_and_merges_two_tech_jobs():
    body = client(FakeStore(MAPPED), api_key="").get("/api/v1/map").json()
    [point] = body["points"]
    assert (point["ref_id"], point["lat"], point["lon"]) == ("j1", 40.123, -100.654)
    assert point["region"] == "North"
    assert [t["tech_id"] for t in point["techs"]] == ["a", "b"]
    assert body["exact"] is False
    [unmapped] = body["unmapped"]
    assert (unmapped["ref_id"], unmapped["address_issue"]) == ("u1", "no_gps")
    assert "lat" not in unmapped


def test_map_serves_exact_coordinates_only_when_configured():
    exact = client(FakeStore(MAPPED), api_key="", exact_coords=True)
    assert exact.get("/api/v1/map").json()["points"][0]["lat"] == 40.123456


def test_map_filters():
    api = client(FakeStore(MAPPED), api_key="")
    assert api.get("/api/v1/map", params={"region": "South"}).json()["points"] == []
    assert api.get("/api/v1/map", params={"kind": "ticket"}).json()["unmapped"] == []
    assert api.get("/api/v1/map", params={"kind": "bogus"}).status_code == 422


def test_diagnostics(api):
    body = api.get("/api/v1/diagnostics").json()
    summary = {c["id"]: c for c in body["summary"]}
    assert summary["work_without_shift"]["count"] == 0
    assert summary["work_address_issue"]["available"] is False
    assert len(body["checks"]) == len(body["summary"])
    tech = api.get("/api/v1/diagnostics", params={"group": "tech_setup"}).json()
    assert {c["group"] for c in tech["checks"]} == {"tech_setup"}
    one = api.get("/api/v1/diagnostics", params={"check": "double_booked"}).json()
    assert [c["id"] for c in one["checks"]] == ["double_booked"]


def test_diagnostics_rejects_unknown_filters(api):
    assert api.get("/api/v1/diagnostics", params={"group": "nope"}).status_code == 422
    assert api.get("/api/v1/diagnostics", params={"check": "nope"}).status_code == 422


TC_BLOCKS = BLOCKS + [
    block("shift_tc", "08:00", "17:00", "t", "North", "TC"),
    block("shift_tc", "08:00", "17:00", "b", "South", "INS"),  # b is a dual tech
    block("ticket", "09:00", "10:00", "t", ref="tk1"),
]


def test_tc_calendar_is_separate_from_install():
    api = client(FakeStore(TC_BLOCKS), api_key="")
    install = api.get("/api/v1/calendar", params={"days": 1}).json()["days"][0]["totals"]
    tc = api.get("/api/v1/calendar", params={"days": 1, "calendar": "tc"}).json()["days"][0]
    assert install["techs_on"] == 2  # a and b only
    assert (tc["totals"]["techs_on"], tc["totals"]["tickets"]) == (2, 1)
    day = api.get("/api/v1/calendar/2026-10-06", params={"calendar": "tc"}).json()
    assert [t["tech_id"] for t in day["techs"]] == ["b", "t"]
    assert api.get("/api/v1/calendar", params={"calendar": "bogus"}).status_code == 422
    techs = {t["tech_id"]: t["calendar"] for t in api.get("/api/v1/filters").json()["techs"]}
    assert techs == {"a": "install", "b": "both", "t": "tc"}
