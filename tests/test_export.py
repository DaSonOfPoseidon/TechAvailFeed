import io
from datetime import date
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from api.capacity import Capacity
from api.export import CAPACITY_FIELDS, workbook


def test_capacity_columns_cover_every_field():
    assert set(CAPACITY_FIELDS) == set(Capacity.__dataclass_fields__)


def test_empty_export_still_opens_with_headers_and_no_tables():
    data = workbook(
        tz=ZoneInfo("America/Chicago"),
        meta={"id": None, "generated_at": None},
        filters={"region": "North"},
        entries=[],
        days=[],
        demand=[],
        schedule=[],
        outcomes=[],
        kpis={"days": []},
        checks=[],
    )
    wb = load_workbook(io.BytesIO(data))
    assert wb["Summary"].max_row == 1
    assert not wb["Summary"].tables
    about = {row[0]: row[1] for row in wb["About"].iter_rows(values_only=True) if row[0]}
    assert about["Filter: region"] == "North"
    assert "Net h" in about


def test_headers_are_purple_with_white_bold_text():
    data = workbook(
        tz=ZoneInfo("America/Chicago"),
        meta={"id": None, "generated_at": None},
        filters={},
        entries=[],
        days=[],
        demand=[],
        schedule=[],
        outcomes=[],
        kpis={"days": []},
        checks=[],
    )
    wb = load_workbook(io.BytesIO(data))
    for ws in wb.worksheets:
        if ws.title == "About":
            continue
        for cell in ws[1]:
            style = (cell.fill.fgColor.rgb, cell.font.color.rgb, cell.font.b)
            assert style == ("007030A0", "00FFFFFF", True)


def test_formula_like_values_are_stored_as_text():
    from feed.diagnostics import Check

    check = Check("tech_no_region", "tech_setup", "No region", "warning", "", rows=[])
    check.rows = [{"tech_id": "=1+1", "tech_name": '=HYPERLINK("http://x","y")'}]
    data = workbook(
        tz=ZoneInfo("America/Chicago"),
        meta={"id": 1, "generated_at": None},
        filters={"region": "=cmd|' /C calc'!A0"},
        entries=[],
        days=[],
        demand=[],
        schedule=[],
        outcomes=[],
        kpis={"days": []},
        checks=[check],
    )
    wb = load_workbook(io.BytesIO(data))
    cells = [cell for ws in wb.worksheets for row in ws.iter_rows() for cell in row]
    assert not [cell.coordinate for cell in cells if cell.data_type == "f"]
    header, row = wb["Diagnostics"].iter_rows(values_only=True)
    assert row[header.index("Tech id")] == "=1+1"
    about = {row[0]: row[1] for row in wb["About"].iter_rows(values_only=True) if row[0]}
    assert about["Filter: region"] == "=cmd|' /C calc'!A0"


def test_a_day_without_a_morning_plan_gets_a_full_width_row():
    kpis = {"days": [{"date": date(2026, 9, 29), "status": "no_morning", "provisional": False}]}
    data = workbook(
        tz=ZoneInfo("America/Chicago"),
        meta={"id": 1, "generated_at": None},
        filters={},
        entries=[],
        days=[],
        demand=[],
        schedule=[],
        outcomes=[],
        kpis=kpis,
        checks=[],
    )
    sheet = load_workbook(io.BytesIO(data))["Outcomes by day"]
    header, row = sheet.iter_rows(values_only=True)
    assert len(row) == len(header)
    assert row[1:4] == ("no_morning", False, None)


def test_diagnostics_sheet_flattens_check_specific_fields_into_detail():
    from feed.diagnostics import Check

    check = Check("tech_no_region", "tech_setup", "No region", "warning", "", rows=[])
    check.rows = [{"tech_id": "a", "tech_name": "A", "days": 3, "first_date": date(2026, 10, 6)}]
    data = workbook(
        tz=ZoneInfo("America/Chicago"),
        meta={"id": 1, "generated_at": None},
        filters={},
        entries=[],
        days=[],
        demand=[],
        schedule=[],
        outcomes=[],
        kpis={"days": []},
        checks=[check],
    )
    wb = load_workbook(io.BytesIO(data))
    header, row = wb["Diagnostics"].iter_rows(values_only=True)
    assert row[header.index("Check")] == "No region"
    assert row[header.index("Tech id")] == "a"
    assert row[header.index("Detail")] == "days: 3; first date: 2026-10-06"
